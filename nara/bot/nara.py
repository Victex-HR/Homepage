#!/usr/bin/env python3
"""NARA 자동화: 공개 사이트 갱신 · 엑셀 생성 · 일정에 맞춘 메일 발송.

GitHub Actions 가 매시간 `python nara.py tick` 을 실행한다.
  1) 오늘(한국시간) 사이트를 아직 갱신하지 않았으면 최근 3일 공고를 다시 받아 nara/data 를 갱신
  2) Supabase 의 발송 일정 중 지금 보낼 차례인 것이 있으면 전날 공고 엑셀을 만들어 수신자에게 메일 발송

수동 실행
  python nara.py update-site               # 최근 3일 수집 + nara/data 갱신
  python nara.py excel --date 20261012     # out/ 에 엑셀 생성
  python nara.py send-test --to a@b.com    # 어제 엑셀을 한 주소로 시험 발송

환경 변수 (GitHub 저장소 Secrets)
  NARA_API_KEY            공공데이터포털 일반 인증키(Decoding)
  SUPABASE_URL            https://xxxx.supabase.co
  SUPABASE_SERVICE_KEY    Supabase service_role 키 (서버 전용, 공개 금지)
  SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD, MAIL_FROM
"""
import argparse
import datetime as dt
import glob
import json
import os
import smtplib
import ssl
import sys
import urllib.parse
import urllib.request
from email.message import EmailMessage

import nara_bids as nb

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.environ.get("NARA_DATA_DIR") or os.path.join(BASE, "..", "data")   # nara-data 브랜치 내용
DAYS = os.path.join(DATA, "days")
PRE = os.path.join(DATA, "pre")    # 사전규격 (접수일별)
OUT = os.path.join(BASE, "out")
KST = dt.timezone(dt.timedelta(hours=9))
KEEP_DAYS = 60          # 사이트에 남겨 둘 공고일 수
TODAY_EVERY_H = 2       # 오늘 공고를 다시 받는 간격
SEND_WINDOW_H = 3       # 예정 시각이 지나고 이 시간 안이면 보냄 (GitHub 실행 지연 대비)
PAGE_URL = os.environ.get("SITE_URL", "")


def now_kst():
    return dt.datetime.now(KST)


def load_cfg():
    """추천 기준. Supabase(관리 화면에서 편집)에 있으면 그것을, 없으면 config.json 을 쓴다."""
    if os.environ.get("SUPABASE_URL") and os.environ.get("SUPABASE_SERVICE_KEY"):
        try:
            rows = sb("GET", "settings?select=value&key=eq.criteria")
            if rows and rows[0].get("value", {}).get("keyword_groups") is not None:
                return rows[0]["value"]
        except Exception as e:  # noqa: BLE001 - 저장소 문제면 파일 기준으로 계속
            print(f"Supabase 기준을 읽지 못해 config.json 을 씀: {e}")
    return json.load(open(os.path.join(BASE, "config.json"), encoding="utf-8"))


# ---------------------------------------------------------------- 수집 · 사이트
def fetch_day(key, day):
    """YYYYMMDD 하루 공고 전체 (1시간 단위로 나눠 받음)."""
    items, total = {}, 0
    for h in range(24):
        got, n = nb.fetch_window(key, f"{day}{h:02d}00", f"{day}{h:02d}59")
        for it in got:
            items[(it.get("bidNtceNo"), it.get("bidNtceOrd"))] = it
        total += n
    print(f"[{day}] {len(items)}건 / 전체 {total}건", flush=True)
    return list(items.values())


def slim(items, cfg, iso):
    out = []
    ref = dt.date.fromisoformat(iso)
    for r in sorted(items, key=lambda r: r.get("bidNtceBgn") or ""):
        row = {f: ("" if r.get(f) is None else str(r.get(f))) for f in nb.FIELDS}
        if row["bidNtceDate"] != iso:
            continue
        s, why = nb.score(row, cfg, ref)
        it = {k: row[k] for k in nb.SLIM}
        it["score"], it["why"] = s, why
        out.append(it)
    return out


def load_meta():
    try:
        return json.load(open(os.path.join(DATA, "index.json"), encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def rebuild_index(cfg, state=None):
    """site/data/days/*.json 으로 index.json · recs.json 을 다시 만든다."""
    files = sorted(glob.glob(os.path.join(DAYS, "*.json")), reverse=True)
    for old in files[KEEP_DAYS:]:
        os.remove(old)
    files = files[:KEEP_DAYS]
    today = now_kst().date()
    index, recs, total = [], {}, 0
    for f in files:
        d = os.path.basename(f)[:8]
        iso = f"{d[:4]}-{d[4:6]}-{d[6:]}"
        rows = json.load(open(f, encoding="utf-8"))
        by_type, n_rec = {}, 0
        for r in rows:
            t = r.get("bsnsDivNm") or "기타"
            by_type[t] = by_type.get(t, 0) + 1
            s, why = nb.score(r, cfg, dt.date.fromisoformat(iso))
            r["score"], r["why"] = s, why
            if s is None:
                continue
            n_rec += 1
            s_now, why_now = nb.score(r, cfg, today)
            if s_now is not None:
                prev = recs.get(r["bidNtceNo"])
                if not prev or r["bidNtceOrd"] > prev["bidNtceOrd"]:
                    recs[r["bidNtceNo"]] = dict(r, score=s_now, why=why_now)
        json.dump(rows, open(f, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
        total += len(rows)
        index.append({"date": iso, "total": len(rows), "byType": by_type, "recommended": n_rec})
    pre_index, oldest = [], index[-1]["date"].replace("-", "") if index else ""
    for f in sorted(glob.glob(os.path.join(PRE, "*.json")), reverse=True):
        d = os.path.basename(f)[:8]
        if d < oldest:      # 입찰공고와 같은 기간만 보관
            os.remove(f)
            continue
        rows = json.load(open(f, encoding="utf-8"))
        for r in rows:
            if "specUrl" not in r:    # 예전 파일: 공고 링크 자리에 규격서 파일 주소가 들어 있었음
                r["specUrl"], r["bidNtceUrl"] = r.get("bidNtceUrl", ""), pre_url(r.get("bidNtceNo"))
            r["score"], r["why"] = nb.score(r, cfg, dt.date(int(d[:4]), int(d[4:6]), int(d[6:])))
        json.dump(rows, open(f, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
        pre_index.append({"date": f"{d[:4]}-{d[4:6]}-{d[6:]}", "total": len(rows),
                          "recommended": sum(r["score"] is not None for r in rows)})
    recs = sorted(recs.values(), key=lambda x: (-x["score"], x["bidClseDate"] or "9999"))
    json.dump(recs, open(os.path.join(DATA, "recs.json"), "w", encoding="utf-8"),
              ensure_ascii=False, separators=(",", ":"))
    write_admin_config()
    prev = load_meta()
    meta = {"generated": now_kst().strftime("%Y-%m-%dT%H:%M"), "sample": False, "public": True,
            "daily_on": prev.get("daily_on"), "today_at": prev.get("today_at"), **(state or {}),
            "company": cfg.get("company_name"), "criteria": cfg, "totalStored": total, "days": index,
            "pre_days": pre_index}
    json.dump(meta, open(os.path.join(DATA, "index.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(f"사이트 데이터: {len(index)}일, 추천 {len(recs)}건")


def write_admin_config():
    """관리 화면이 읽는 공개 설정 (Supabase 주소와 anon 키는 공개해도 되는 값)."""
    url, anon = os.environ.get("SUPABASE_URL"), os.environ.get("SUPABASE_ANON_KEY")
    if url and anon:
        json.dump({"supabaseUrl": url, "supabaseAnonKey": anon},
                  open(os.path.join(DATA, "admin-config.json"), "w"))


# ---------------------------------------------------------------- 사전규격
PRE_URL = "https://apis.data.go.kr/1230000/ao/HrcspSsstndrdInfoService/"
PRE_OPS = {"getPublicPrcureThngInfoThng": "물품", "getPublicPrcureThngInfoServc": "용역",
           "getPublicPrcureThngInfoCnstwk": "공사", "getPublicPrcureThngInfoFrgcpt": "외자"}


def pre_page(key, op, begin, end, page):
    q = urllib.parse.urlencode({"ServiceKey": key, "numOfRows": nb.PAGE_SIZE, "pageNo": page, "type": "json",
                                "inqryDiv": 1, "inqryBgnDt": begin, "inqryEndDt": end})
    for attempt in range(8):
        try:
            with urllib.request.urlopen(f"{PRE_URL}{op}?{q}", timeout=30) as r:
                data = json.loads(r.read().decode("utf-8"))
            break
        except (OSError, ValueError) as e:
            if attempt == 7:
                raise RuntimeError(f"사전규격 API 호출 실패: {e}")
            nb.time.sleep(min(30, 2 ** (attempt + 1)))
    if "OpenAPI_ServiceResponse" in data:   # 활용신청이 안 된 키 등
        h = data["OpenAPI_ServiceResponse"].get("cmmMsgHeader", {})
        raise RuntimeError(f"사전규격 API: {h.get('returnAuthMsg') or h.get('errMsg')}")
    resp = data.get("response", data)
    hdr = resp.get("header", {})
    if str(hdr.get("resultCode", "00")) not in ("00", "0"):
        raise RuntimeError(f"사전규격 API 오류 {hdr.get('resultCode')}: {hdr.get('resultMsg')}")
    body = resp.get("body", {})
    items = body.get("items") or []
    if isinstance(items, dict):
        items = items.get("item", [])
    if isinstance(items, dict):
        items = [items]
    return items, int(body.get("totalCount") or 0)


def pre_url(no):
    return f"https://www.g2b.go.kr/link/PRVA004_02/single/?bfSpecRegNo={no}" if no else ""


def pre_slim(r, kind):
    """사전규격 한 건을 공고 목록과 같은 모양으로 바꾼다 (점수 계산·화면 재사용)."""
    g = lambda *ks: next((str(r[k]) for k in ks if r.get(k) not in (None, "")), "")
    rcpt, close = g("rcptDt", "rgstDt"), g("opninRgstClseDt")
    no = g("bfSpecRgstNo")
    return {
        "bidNtceNo": no, "bidNtceOrd": "000", "bidNtceNm": g("prdctClsfcNoNm", "bsnsNm", "prdctNm"),
        "bidNtceSttusNm": "", "bidNtceDate": rcpt[:10], "bidNtceBgn": rcpt[11:16],
        "bsnsDivNm": kind, "preKind": g("bsnsDivNm"),   # 업무구분 필터와 맞추려고 물품/용역/공사/외자로 "cntrctCnclsMthdNm": "", "bidwinrDcsnMthdNm": "",
        "ntceInsttNm": g("orderInsttNm", "ntceInsttNm"), "dmndInsttNm": g("rlDminsttNm", "dminsttNm"),
        "bidPrtcptQlfctRgstClseDate": "", "bidPrtcptQlfctRgstClseTm": "",
        "bidClseDate": close[:10], "bidClseTm": close[11:16], "opengDate": "", "opengTm": "",
        "asignBdgtAmt": g("asignBdgtAmt"), "presmptPrce": "", "rgnLmtYn": "", "prtcptPsblRgnNm": "",
        "indstrytyLmtYn": "", "bidprcPsblIndstrytyNm": "",
        "bidNtceUrl": pre_url(no),            # 나라장터 사전규격 상세 화면
        "specUrl": g("specDocFileUrl1"),      # 규격서 첫 파일 (나라장터 내려받기 주소)
        "pre": 1, "bidNtceNoList": g("bidNtceNoList"), "ofcl": " ".join(x for x in (g("ofclNm"), g("ofclTelNo")) if x),
    }


def update_pre(days):
    """사전규격을 접수일 기준으로 받아 data/pre/YYYYMMDD.json 에 저장. 실패해도 입찰공고 수집은 계속."""
    key = os.environ.get("NARA_API_KEY")
    cfg = load_cfg()
    os.makedirs(PRE, exist_ok=True)
    try:
        for d in days:
            got = {}
            for op, kind in PRE_OPS.items():
                begin, end = f"{d:%Y%m%d}0000", f"{d:%Y%m%d}2359"
                items, total = pre_page(key, op, begin, end, 1)
                for p in range(2, (total + nb.PAGE_SIZE - 1) // nb.PAGE_SIZE + 1):
                    items += pre_page(key, op, begin, end, p)[0]
                for r in items:     # 조회일(등록 기준)에 묶음. 화면에는 접수일시를 보여 줌
                    row = pre_slim(r, kind)
                    got[row["bidNtceNo"]] = row
            rows = sorted(got.values(), key=lambda r: r["bidNtceBgn"])
            for r in rows:
                r["score"], r["why"] = nb.score(r, cfg, d)
            json.dump(rows, open(os.path.join(PRE, f"{d:%Y%m%d}.json"), "w", encoding="utf-8"),
                      ensure_ascii=False, separators=(",", ":"))
            print(f"[사전규격 {d}] {len(rows)}건", flush=True)
        return None
    except RuntimeError as e:
        print(e, flush=True)
        return str(e)


def update_site(days_back=3, include_today=True, only_today=False):
    """지난 days_back 일(어제부터)과 오늘 공고를 다시 받아 저장. only_today 면 오늘만."""
    key = os.environ.get("NARA_API_KEY")
    if not key:
        raise SystemExit("NARA_API_KEY 가 없습니다.")
    cfg = load_cfg()
    os.makedirs(DAYS, exist_ok=True)
    today = now_kst().date()
    offsets = [0] if only_today else list(range(days_back, 0, -1)) + ([0] if include_today else [])
    pre_err = update_pre([today - dt.timedelta(days=i) for i in offsets])
    for i in offsets:
        d = today - dt.timedelta(days=i)
        rows = slim(fetch_day(key, d.strftime("%Y%m%d")), cfg, d.isoformat())
        json.dump(rows, open(os.path.join(DAYS, d.strftime("%Y%m%d") + ".json"), "w", encoding="utf-8"),
                  ensure_ascii=False, separators=(",", ":"))
    now = now_kst()
    state = {"today_at": now.strftime("%Y-%m-%dT%H:%M")} if 0 in offsets else {}
    state["pre_error"] = pre_err
    if not only_today:
        state["daily_on"] = today.isoformat()
    rebuild_index(cfg, state)


def update_today_new():
    """오늘 공고 중 이미 가진 마지막 공고 시각 이후만 받아 덧붙인다 (몇 초~수십 초)."""
    key = os.environ.get("NARA_API_KEY")
    if not key:
        raise SystemExit("NARA_API_KEY 가 없습니다.")
    cfg, now = load_cfg(), now_kst()
    d = now.date()
    path = os.path.join(DAYS, d.strftime("%Y%m%d") + ".json")
    try:
        rows = json.load(open(path, encoding="utf-8"))
    except (OSError, ValueError):
        rows = []
    last = max((r.get("bidNtceBgn") or "00:00" for r in rows), default="00:00")
    have = {(r["bidNtceNo"], r["bidNtceOrd"]) for r in rows}
    got = []
    for h in range(int(last[:2]), now.hour + 1):
        items, _ = nb.fetch_window(key, f"{d:%Y%m%d}{h:02d}00", f"{d:%Y%m%d}{h:02d}59")
        got.extend(items)
    new = [r for r in slim(got, cfg, d.isoformat()) if (r["bidNtceNo"], r["bidNtceOrd"]) not in have]
    rows = sorted(rows + new, key=lambda r: r.get("bidNtceBgn") or "")
    os.makedirs(DAYS, exist_ok=True)
    json.dump(rows, open(path, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    print(f"[{d}] {last} 이후 새 공고 {len(new)}건", flush=True)
    pre_err = update_pre([d])               # 사전규격은 하루치가 적어 매번 오늘 전체를 다시 받음
    rebuild_index(cfg, {"new_at": now.strftime("%Y-%m-%dT%H:%M"), "pre_error": pre_err})
    return bool(new)


# ---------------------------------------------------------------- 엑셀
COLS = [("추천점수", "score", 9), ("공고명", "bidNtceNm", 60), ("업무구분", "bsnsDivNm", 9),
        ("공고기관", "ntceInsttNm", 26), ("수요기관", "dmndInsttNm", 26),
        ("배정예산(원)", "asignBdgtAmt", 15), ("추정가격(원)", "presmptPrce", 15),
        ("입찰마감", "close", 17), ("개찰", "open", 17), ("계약방법", "cntrctCnclsMthdNm", 14),
        ("낙찰방법", "bidwinrDcsnMthdNm", 18), ("지역제한", "rgn", 18), ("업종제한", "inds", 30),
        ("추천 사유", "why", 40), ("공고번호", "no", 18), ("공고일", "bidNtceDate", 11),
        ("공고 보기", "bidNtceUrl", 12)]


def _cell(r, k):
    if k == "close":
        return f'{r.get("bidClseDate", "")} {r.get("bidClseTm", "")}'.strip()
    if k == "open":
        return f'{r.get("opengDate", "")} {r.get("opengTm", "")}'.strip()
    if k == "rgn":
        return (r.get("prtcptPsblRgnNm") or "있음") if r.get("rgnLmtYn") == "Y" else ""
    if k == "inds":
        return (r.get("bidprcPsblIndstrytyNm") or "있음") if r.get("indstrytyLmtYn") == "Y" else ""
    if k == "why":
        return ", ".join(r.get("why") or []) if r.get("score") is not None else ""
    if k == "no":
        return f'{r.get("bidNtceNo", "")}-{r.get("bidNtceOrd", "")}'
    if k in ("asignBdgtAmt", "presmptPrce"):
        return nb.to_int(r.get(k)) or None
    if k == "score":
        return r.get("score")
    return r.get(k) or ""


def _sheet(wb, title, rows, first=False):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append([c[0] for c in COLS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B6B63")
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([_cell(r, k) for _, k, _ in COLS])
        url = r.get("bidNtceUrl")
        if url:
            link = ws.cell(row=ws.max_row, column=len(COLS))
            link.value, link.hyperlink, link.font = "공고 보기", url, Font(color="0563C1", underline="single")
    for i, (_, k, w) in enumerate(COLS, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        if k in ("asignBdgtAmt", "presmptPrce"):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = "#,##0"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build_excel(day):
    """day(YYYYMMDD) 공고 엑셀. (경로, 요약) 반환."""
    from openpyxl import Workbook
    f = os.path.join(DAYS, f"{day}.json")
    rows = json.load(open(f, encoding="utf-8")) if os.path.exists(f) else []
    recs_day = sorted([r for r in rows if r.get("score") is not None],
                      key=lambda r: (-r["score"], r.get("bidClseDate") or "9"))
    try:
        open_recs = json.load(open(os.path.join(DATA, "recs.json"), encoding="utf-8"))
    except (OSError, ValueError):
        open_recs = []
    wb = Workbook()
    _sheet(wb, "추천 공고", recs_day, first=True)
    _sheet(wb, "진행 중 추천 전체", open_recs)
    _sheet(wb, "전체 공고", sorted(rows, key=lambda r: -(r.get("score") or -1)))
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, f"나라장터_입찰공고_{day}.xlsx")
    wb.save(path)
    return path, {"day": day, "total": len(rows), "recs": recs_day, "open": len(open_recs)}


# ---------------------------------------------------------------- 메일
def send_mail(to, path, summary):
    host, user = os.environ.get("SMTP_HOST"), os.environ.get("SMTP_USER")
    if not host or not to:
        raise SystemExit("SMTP_HOST 또는 수신자가 없습니다.")
    port = int(os.environ.get("SMTP_PORT") or 465)
    d = summary["day"]
    iso = f"{d[:4]}-{d[4:6]}-{d[6:]}"
    msg = EmailMessage()
    msg["Subject"] = f"[나라장터] {iso} 입찰공고 {summary['total']:,}건 · 추천 {len(summary['recs'])}건"
    msg["From"] = os.environ.get("MAIL_FROM") or user
    msg["To"] = ", ".join(to)
    top = summary["recs"][:10]
    lines = [f"{iso} 나라장터 입찰공고 {summary['total']:,}건 중 추천 {len(summary['recs'])}건입니다.",
             f"마감 전인 추천 공고는 모두 {summary['open']}건이며 첨부 엑셀에 정리했습니다.", ""]
    lines += [f"{i}. [{r['score']}점] {r['bidNtceNm']} / {r['ntceInsttNm']} / 마감 {r['bidClseDate']} {r['bidClseTm']}"
              for i, r in enumerate(top, 1)]
    if PAGE_URL:
        lines += ["", f"웹에서 보기: {PAGE_URL}"]
    msg.set_content("\n".join(lines))
    esc = lambda s: str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    trs = "".join(
        f'<tr><td style="padding:6px 8px;color:#0b6b63;font-weight:600">{r["score"]}</td>'
        f'<td style="padding:6px 8px"><a href="{esc(r["bidNtceUrl"])}">{esc(r["bidNtceNm"])}</a><br>'
        f'<span style="color:#5b6875;font-size:12px">{esc(r["ntceInsttNm"])} · 마감 {esc(r["bidClseDate"])} {esc(r["bidClseTm"])}</span></td></tr>'
        for r in top)
    html = (f'<div style="font-family:Malgun Gothic,Apple SD Gothic Neo,sans-serif;font-size:14px;color:#16202a">'
            f'<p>{iso} 나라장터 입찰공고 <b>{summary["total"]:,}건</b> 중 추천 <b>{len(summary["recs"])}건</b>입니다. '
            f'마감 전인 추천 공고 {summary["open"]}건은 첨부 엑셀에 정리했습니다.</p>'
            f'<table style="border-collapse:collapse">{trs}</table>'
            + (f'<p><a href="{PAGE_URL}">웹에서 보기</a></p>' if PAGE_URL else "") + "</div>")
    msg.add_alternative(html, subtype="html")
    with open(path, "rb") as fh:
        msg.add_attachment(fh.read(), maintype="application",
                           subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           filename=os.path.basename(path))
    ctx = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=ctx, timeout=60) as s:
            s.login(user, os.environ["SMTP_PASSWORD"]); s.send_message(msg)
    else:
        with smtplib.SMTP(host, port, timeout=60) as s:
            s.starttls(context=ctx); s.login(user, os.environ["SMTP_PASSWORD"]); s.send_message(msg)


# ---------------------------------------------------------------- Supabase
def sb(method, path, body=None):
    url = os.environ["SUPABASE_URL"].rstrip("/") + "/rest/v1/" + path
    key = os.environ["SUPABASE_SERVICE_KEY"]
    req = urllib.request.Request(url, method=method, data=json.dumps(body).encode() if body is not None else None,
                                 headers={"apikey": key, "Authorization": f"Bearer {key}",
                                          "Content-Type": "application/json", "Prefer": "return=minimal"})
    with urllib.request.urlopen(req, timeout=30) as r:
        raw = r.read()
    return json.loads(raw) if raw else None


def due_schedules(now):
    today = now.date()
    out = []
    for s in sb("GET", "schedules?select=*&active=eq.true") or []:
        if s.get("last_sent_on") == today.isoformat():
            continue
        if now.isoweekday() not in (s.get("weekdays") or [1, 2, 3, 4, 5, 6, 7]):
            continue
        hh, mm = map(int, s["send_time"].split(":")[:2])
        at = dt.datetime.combine(today, dt.time(hh, mm), KST)
        if at <= now < at + dt.timedelta(hours=SEND_WINDOW_H):
            out.append(s)
    return out


# Supabase 설정 전에는 MAIL_TO(쉼표 구분)로 매일 MAIL_TIME(기본 00:00)에 보냄. 기록은 nara-data 의 mail-state.json
STATE = os.path.join(DATA, "mail-state.json")


def fallback_due(now):
    if not os.environ.get("MAIL_TO"):
        return []
    try:
        last = json.load(open(STATE)).get("last_sent_on")
    except (OSError, ValueError):
        last = None
    hh, mm = map(int, (os.environ.get("MAIL_TIME") or "00:00").split(":")[:2])
    at = dt.datetime.combine(now.date(), dt.time(hh, mm), KST)
    if last != now.date().isoformat() and at <= now < at + dt.timedelta(hours=SEND_WINDOW_H):
        return [{"id": None, "name": "기본 일정", "fallback": True}]
    return []


def tick():
    now = now_kst()
    has_sb = bool(os.environ.get("SUPABASE_URL") and os.environ.get("SUPABASE_SERVICE_KEY"))
    due = due_schedules(now) if has_sb else fallback_due(now)
    changed = False
    meta = load_meta()
    if meta.get("daily_on") != now.date().isoformat():
        update_site()                       # 하루 한 번: 사흘 전~어제 다시 받기 + 오늘
        changed = True
    else:
        last = meta.get("today_at") or ""
        if not last.startswith(now.date().isoformat()) or \
                now - dt.datetime.fromisoformat(last).replace(tzinfo=KST) >= dt.timedelta(hours=TODAY_EVERY_H):
            update_site(only_today=True)    # 2시간마다 오늘 공고 전체 (정정·누락 보정)
            changed = True
        else:
            update_today_new()              # 그 사이에는 매번 새 공고만 덧붙임
            changed = True
    # 관리 화면에서 기준을 바꾸면 수집 없이 점수만 다시 매겨 공개 화면에 반영
    if not changed and meta and load_cfg() != meta.get("criteria"):
        rebuild_index(load_cfg())
        changed = True
    cfg_path = os.path.join(DATA, "admin-config.json")
    before = open(cfg_path).read() if os.path.exists(cfg_path) else ""
    write_admin_config()
    if os.path.exists(cfg_path) and open(cfg_path).read() != before:
        changed = True
    if due:
        if has_sb:
            rcpts = [r["email"] for r in sb("GET", "recipients?select=email&active=eq.true") or []]
        else:
            rcpts = [x.strip() for x in os.environ["MAIL_TO"].split(",") if x.strip()]
        day = (now.date() - dt.timedelta(days=1)).strftime("%Y%m%d")
        path, summary = build_excel(day)
        for s in due:
            status, detail = "sent", f"{summary['total']}건 · 추천 {len(summary['recs'])}건"
            try:
                if not rcpts:
                    raise RuntimeError("활성 수신자가 없습니다")
                send_mail(rcpts, path, summary)
            except Exception as e:  # noqa: BLE001 - 기록 후 다음 일정 진행
                status, detail = "failed", f"{type(e).__name__}: {e}"[:500]
            if s.get("fallback"):
                if status == "sent":
                    json.dump({"last_sent_on": now.date().isoformat()}, open(STATE, "w"))
                    changed = True
            else:
                sb("POST", "send_log", {"schedule_id": s["id"], "recipients": rcpts, "status": status,
                                        "detail": detail, "report_date": f"{day[:4]}-{day[4:6]}-{day[6:]}"})
                if status == "sent":
                    sb("PATCH", f"schedules?id=eq.{s['id']}", {"last_sent_on": now.date().isoformat()})
            print(f"일정 '{s['name']}': {status} {detail}")
    # GitHub Actions 다음 단계에서 사이트 반영 여부 판단
    if os.environ.get("GITHUB_OUTPUT"):
        with open(os.environ["GITHUB_OUTPUT"], "a") as fh:
            fh.write(f"changed={'true' if changed else 'false'}\n")
    print(f"{now:%Y-%m-%d %H:%M} 사이트 갱신 {'함' if changed else '안 함'}, 발송 일정 {len(due)}건")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["tick", "update-site", "rebuild", "excel", "send-test"])
    ap.add_argument("--date", help="YYYYMMDD (기본: 어제)")
    ap.add_argument("--days-back", type=int, default=3)
    ap.add_argument("--to", action="append")
    a = ap.parse_args()
    day = a.date or (now_kst().date() - dt.timedelta(days=1)).strftime("%Y%m%d")
    if a.cmd == "tick":
        tick()
    elif a.cmd == "update-site":
        update_site(a.days_back)
    elif a.cmd == "rebuild":
        rebuild_index(load_cfg())
    elif a.cmd == "excel":
        print(build_excel(day)[0])
    else:
        path, summary = build_excel(day)
        send_mail(a.to, path, summary)
        print(f"시험 발송: {', '.join(a.to)}")


if __name__ == "__main__":
    sys.exit(main())
